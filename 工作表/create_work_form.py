from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt
import datetime
import calendar

day_to_chinese = {1: '一', 2: '二', 3: '三', 4: '四', 5: '五'}


def convert_day_to_chinese(day):
    return day_to_chinese.get(day, 'None')


def iter_month(month):  # 無考慮閏年 產生那一個月的迭代器
    beg_date = datetime.date(datetime.datetime.now().date().year, month, 1)
    end_date = beg_date + datetime.timedelta(calendar.mdays[month])
    while beg_date != end_date:
        yield beg_date
        beg_date += datetime.timedelta(1)
def process_chinese_setting(run):  # 使run正常顯示中文(標楷體)
    run.font.name = u'標楷體'
    r = run._element.rPr.rFonts  # 中文特有的處理
    r.set(qn("w:eastAsia"), "標楷體")

def table_run_add_and_set(table, row, col, text):  # add run and set some value in table
    run = table.cell(row, col).paragraphs[0].add_run(text)
    table.cell(row, col).paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER  # 水平置中
    process_chinese_setting(run)


def text_run_add_and_set(docx, idx, pre, hour, suf):  # add run and set some value in text
    docx.paragraphs[idx].text = ''
    run = docx.paragraphs[idx].add_run(pre)
    process_chinese_setting(run)
    run.font.size = Pt(14)
    run.bold = True
    run = docx.paragraphs[idx].add_run(str(hour))
    process_chinese_setting(run)
    run.font.size = Pt(14)
    run.bold = True
    run.underline = True
    run = docx.paragraphs[idx].add_run(suf)
    process_chinese_setting(run)
    run.font.size = Pt(14)
    run.bold = True


def generate_table(table, hour, days, month):  # hour is integer,days is a list contain integer,month is int
    nowrow = 1
    for date in iter_month(month):
        if date.isoweekday() in days:
            if hour > 7:
                table_run_add_and_set(table, nowrow, 0, f'{date.year - 1911}/{date.month}/{date.day}')
                table_run_add_and_set(table, nowrow, 1, convert_day_to_chinese(date.isoweekday()))
                table_run_add_and_set(table, nowrow, 2, '13:00~21:00\n(17:00-18:00休息)')
                table_run_add_and_set(table, nowrow, 3, str(7))
                table_run_add_and_set(table, nowrow, 4, '資料整理')
            else:
                table_run_add_and_set(table, nowrow, 0, f'{date.year - 1911}/{date.month}/{date.day}')
                table_run_add_and_set(table, nowrow, 1, convert_day_to_chinese(date.isoweekday()))
                if hour > 4:
                    table_run_add_and_set(table, nowrow, 2, '13:00~' + str(13 + hour + 1) + ':00\n(17:00-18:00休息)')
                else:
                    table_run_add_and_set(table, nowrow, 2, '13:00~' + str(13 + hour) + ':00')
                table_run_add_and_set(table, nowrow, 3, str(hour))
                table_run_add_and_set(table, nowrow, 4, '資料整理')
                return True
            hour -= 7
            nowrow += 1
    return False


docx = Document("work.docx")
table = docx.tables[0]
available_day = [3, 5]  # list(map(int,input("請輸入可行的工作日: ").split()))
hour = int(input("請輸入工時: "))
month = int(input("請輸入月份: "))
if generate_table(table, hour, available_day, month):
    print('產生word完成。')
    text_run_add_and_set(docx, 2, '*正常工作時數：', f'      {hour}     ', '小時')
    text_run_add_and_set(docx, 3, '*薪資小計：', f'      {hour*160}      ', '元(A)')
    text_run_add_and_set(docx, 5, '*合計應領薪資：', f'      {hour*160}     ', '元(A+B) /　工讀生指導人：______________')
    docx.save(str(month) + '月份工作表.docx')
else:
    print('工作日不夠，產生word失敗。')
